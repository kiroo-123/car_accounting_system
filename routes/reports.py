# -*- coding: utf-8 -*-
"""
routes/reports.py
==================
كل التقارير: الرحلات، الأرباح، العربيات، العملاء، أصحاب العربيات،
السواقين، السولار، الصيانة، المصروفات، التحصيل.
كل التقارير (ما عدا التحصيل) بتتعرض بجدول عام واحد (columns/rows/totals)
مبني على نفس الـ macro المستخدم في باقي الشاشات.
"""
import io
from flask import Blueprint, render_template, request, send_file
from openpyxl import Workbook

import data_access as da

reports_bp = Blueprint("reports", __name__, url_prefix="/reports")

REPORT_LIST = [
    ("trips", "تقرير الرحلات", "كل الرحلات بالتفصيل مع إمكانية الفلترة بالتاريخ"),
    ("profit", "تقرير الأرباح", "الإيرادات والمصروفات وصافي الربح شهر بشهر"),
    ("vehicles", "تقرير العربيات", "ربح كل عربية على حدة"),
    ("customers", "تقرير العملاء", "مستحقات وتحصيلات كل عميل"),
    ("owners", "تقرير أصحاب العربيات", "مستحقات ومدفوعات كل صاحب عربية"),
    ("drivers", "تقرير السواقين", "مستحقات ومدفوعات كل سواق"),
    ("fuel", "تقرير السولار", "استهلاك وتكلفة السولار لكل عربية"),
    ("maintenance", "تقرير الصيانة", "كل مصاريف الصيانة بالتفصيل"),
    ("expenses", "تقرير المصروفات", "فين رايحة فلوس الشركة بالتفصيل"),
    ("collection", "تقرير التحصيل", "موقف التحصيل من العملاء والدفع لأصحاب العربيات والسواقين"),
]


def _sum(rows, key):
    return round(sum((r.get(key) or 0) for r in rows), 2)


def _build(name):
    """يرجع (title, columns, rows, totals, stat_cards) لأي تقرير غير التحصيل."""
    if name == "trips":
        filters = {k: request.args.get(k, "").strip() for k in ("date_from", "date_to")}
        filters = {k: v for k, v in filters.items() if v}
        rows = da.list_trips(filters)
        columns = [
            {"key": "id", "label": "رقم الرحلة"},
            {"key": "date", "label": "التاريخ"},
            {"key": "customer_name", "label": "العميل"},
            {"key": "vehicle_name", "label": "العربية"},
            {"key": "driver_name", "label": "السواق"},
            {"key": "revenue", "label": "سعر الرحلة", "format": "currency"},
            {"key": "fuel_cost", "label": "السولار", "format": "currency"},
            {"key": "toll_cost", "label": "الكارتة", "format": "currency"},
            {"key": "driver_fee", "label": "أجر السواق", "format": "currency"},
            {"key": "owner_fee", "label": "أجرة صاحب العربية", "format": "currency"},
            {"key": "net_profit", "label": "صافي الربح", "format": "currency"},
        ]
        totals = {k: _sum(rows, k) for k in ("revenue", "fuel_cost", "toll_cost", "driver_fee", "owner_fee", "net_profit")}
        return "تقرير الرحلات", columns, rows, totals, None

    if name == "vehicles":
        rows = da.get_vehicles_report()
        columns = [
            {"key": "name", "label": "العربية"},
            {"key": "ownership_label", "label": "نوع الملكية"},
            {"key": "vehicle_type", "label": "النوع"},
            {"key": "trip_count", "label": "عدد الرحلات", "format": "number"},
            {"key": "revenue", "label": "الإيرادات", "format": "currency"},
            {"key": "fuel_cost", "label": "السولار", "format": "currency"},
            {"key": "toll_cost", "label": "الكارتة", "format": "currency"},
            {"key": "driver_cost", "label": "السواقين", "format": "currency"},
            {"key": "owner_cost", "label": "أصحاب العربيات", "format": "currency"},
            {"key": "maintenance_cost", "label": "الصيانة", "format": "currency"},
            {"key": "total_cost", "label": "إجمالي التكلفة", "format": "currency"},
            {"key": "net_profit", "label": "صافي الربح", "format": "currency"},
        ]
        totals = {k: _sum(rows, k) for k in ("trip_count", "revenue", "fuel_cost", "toll_cost", "driver_cost", "owner_cost", "maintenance_cost", "total_cost", "net_profit")}
        return "تقرير العربيات", columns, rows, totals, None

    if name in ("customers", "owners", "drivers"):
        fn = {"customers": da.get_customers_report, "owners": da.get_owners_report, "drivers": da.get_drivers_report}[name]
        label = {"customers": "العميل", "owners": "صاحب العربية", "drivers": "السواق"}[name]
        title = {"customers": "تقرير العملاء", "owners": "تقرير أصحاب العربيات", "drivers": "تقرير السواقين"}[name]
        rows = fn()
        columns = [
            {"key": "name", "label": label},
            {"key": "trip_count", "label": "عدد الرحلات", "format": "number"},
            {"key": "total_dues", "label": "إجمالي المستحق", "format": "currency"},
            {"key": "collected", "label": "المحصّل / المدفوع", "format": "currency"},
            {"key": "remaining", "label": "المتبقي", "format": "currency"},
        ]
        totals = {k: _sum(rows, k) for k in ("trip_count", "total_dues", "collected", "remaining")}
        return title, columns, rows, totals, None

    if name == "expenses":
        data = da.get_expenses_report(
            date_from=request.args.get("date_from") or None,
            date_to=request.args.get("date_to") or None,
        )
        columns = [{"key": "category", "label": "البند"}, {"key": "amount", "label": "المبلغ", "format": "currency"}]
        totals = {"amount": data["grand_total"]}
        return "تقرير المصروفات", columns, data["rows"], totals, None

    if name == "maintenance":
        rows = da.get_maintenance_report()
        columns = [
            {"key": "date", "label": "التاريخ"},
            {"key": "vehicle_name", "label": "العربية"},
            {"key": "expense_type", "label": "نوع المصروف"},
            {"key": "description", "label": "الوصف"},
            {"key": "amount", "label": "المبلغ", "format": "currency"},
        ]
        totals = {"amount": _sum(rows, "amount")}
        return "تقرير الصيانة", columns, rows, totals, None

    if name == "fuel":
        data = da.get_fuel_report()
        columns = [
            {"key": "name", "label": "العربية"},
            {"key": "trip_count", "label": "عدد الرحلات", "format": "number"},
            {"key": "liters", "label": "اللترات", "format": "number"},
            {"key": "cost", "label": "التكلفة", "format": "currency"},
        ]
        rows = data["by_vehicle"]
        totals = {"trip_count": sum(r["trip_count"] for r in rows), "liters": data["total_liters"], "cost": data["total_cost"]}
        stat_cards = [
            {"label": "إجمالي لترات السولار", "text": f"{data['total_liters']:,.0f} لتر"},
            {"label": "إجمالي تكلفة السولار", "amount": data["total_cost"]},
            {"label": "متوسط سعر اللتر", "amount": data["avg_price"]},
        ]
        return "تقرير السولار", columns, rows, totals, stat_cards

    if name == "profit":
        rows = da.get_profit_report()
        columns = [
            {"key": "month", "label": "الشهر"},
            {"key": "trip_count", "label": "عدد الرحلات", "format": "number"},
            {"key": "total_revenue", "label": "الإيرادات", "format": "currency"},
            {"key": "total_expenses", "label": "المصروفات", "format": "currency"},
            {"key": "net_profit", "label": "صافي الربح", "format": "currency"},
            {"key": "collected", "label": "المحصّل بالشهر", "format": "currency"},
        ]
        return "تقرير الأرباح الشهري", columns, rows, None, None

    return None, None, None, None, None


@reports_bp.route("/")
def index():
    return render_template("reports/index.html", reports=REPORT_LIST)


@reports_bp.route("/<name>")
def view_report(name):
    if name == "collection":
        report = da.get_collection_report()
        return render_template("reports/collection.html", report=report)

    title, columns, rows, totals, stat_cards = _build(name)
    if columns is None:
        return render_template("error.html", code=404, message="التقرير غير موجود"), 404
    return render_template(
        "reports/table.html", title=title, name=name, columns=columns, rows=rows,
        totals=totals, stat_cards=stat_cards, filters=request.args,
    )


@reports_bp.route("/<name>/export")
def export_report(name):
    wb = Workbook()
    wb.remove(wb.active)

    if name == "collection":
        report = da.get_collection_report()
        sheets = [
            ("عملاء", report["customers"], "العميل"),
            ("أصحاب عربيات", report["owners"], "صاحب العربية"),
            ("سواقين", report["drivers"], "السواق"),
        ]
        for sheet_title, rows, label in sheets:
            ws = wb.create_sheet(sheet_title)
            ws.append([label, "عدد الرحلات", "إجمالي المستحق", "المحصّل / المدفوع", "المتبقي"])
            for r in rows:
                ws.append([r["name"], r["trip_count"], r["total_dues"], r["collected"], r["remaining"]])
    else:
        title, columns, rows, totals, _ = _build(name)
        if columns is None:
            return render_template("error.html", code=404, message="التقرير غير موجود"), 404
        ws = wb.create_sheet((title or "تقرير")[:31])
        ws.append([c["label"] for c in columns])
        for row in rows:
            ws.append([row.get(c["key"], "") for c in columns])
        if totals:
            ws.append([totals.get(c["key"], "" if i > 0 else "الإجمالي") for i, c in enumerate(columns)])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, download_name=f"{name}_report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
