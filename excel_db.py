# -*- coding: utf-8 -*-
"""
excel_db.py
===========
طبقة الوصول لملف الإكسل (قاعدة البيانات الأساسية للنظام).

هذا الملف مسؤول فقط عن "كيفية" التعامل مع ملف الإكسل بأمان:
- تعريف هيكل الشيتات (SHEETS).
- فتح/حفظ الملف بشكل آمن (Locking + Atomic writes).
- عمليات عامة (CRUD) تعمل على أي شيت بالاعتماد على تعريف الأعمدة.
- النسخ الاحتياطي.

ملاحظة مهمة عن الأعمدة:
كل عمود معرّف بمفتاح إنجليزي (يُستخدم داخليًا في الكود) وعنوان عربي
(يظهر في صف العناوين داخل ملف الإكسل نفسه عند فتحه في Excel).
القراءة/الكتابة تتم دائمًا بالموضع (index) مش بالبحث عن نص العنوان،
عشان نضمن استقرار الكود حتى لو المستخدم غيّر شكل العناوين يدويًا.
"""

import os
import shutil
from pathlib import Path
from datetime import datetime
from contextlib import contextmanager

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from filelock import FileLock

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
BACKUP_DIR = BASE_DIR / "backups"
DB_PATH = DATA_DIR / "company_data.xlsx"
LOCK_PATH = DATA_DIR / "company_data.lock"

DATA_DIR.mkdir(exist_ok=True, parents=True)
BACKUP_DIR.mkdir(exist_ok=True, parents=True)

# ---------------------------------------------------------------------------
# تعريف هيكل الشيتات: كل شيت = قائمة من (المفتاح بالإنجليزي، العنوان بالعربي)
# ---------------------------------------------------------------------------
SHEETS = {
    "Settings": [
        ("id", "المعرف"),
        ("category", "النوع"),          # general / toll_route / expense_type / general_expense_type
        ("key", "المفتاح"),
        ("value", "القيمة"),
        ("notes", "ملاحظات"),
    ],
    "Customers": [
        ("id", "المعرف"),
        ("name", "اسم العميل"),
        ("phone", "رقم الهاتف"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Vehicles": [
        ("id", "المعرف"),
        ("plate_number", "رقم اللوحة"),
        ("code_name", "اسم / كود العربية"),
        ("vehicle_type", "نوع العربية"),
        ("capacity", "السعة / الإمكانيات"),
        ("ownership_type", "نوع الملكية"),      # company / free
        ("owner_id", "معرف صاحب العربية"),
        ("status", "حالة العربية"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Vehicle_Owners": [
        ("id", "المعرف"),
        ("name", "اسم صاحب العربية"),
        ("phone", "رقم الهاتف"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Drivers": [
        ("id", "المعرف"),
        ("name", "اسم السواق"),
        ("phone", "رقم الهاتف"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Trips": [
        ("id", "رقم الرحلة"),
        ("date", "التاريخ"),
        ("customer_id", "معرف العميل"),
        ("vehicle_id", "معرف العربية"),
        ("driver_id", "معرف السواق"),
        ("route_name", "المسار"),
        ("trip_price", "سعر الرحلة من العميل"),
        ("owner_fee", "أجرة صاحب العربية"),
        ("driver_fee", "أجر السواق"),
        ("fuel_liters", "لترات السولار"),
        ("fuel_price_per_liter", "سعر لتر السولار"),
        ("fuel_cost", "تكلفة السولار (محسوبة)"),
        ("toll_cost", "الكارتة"),
        ("commission", "العمولة (معلوماتية)"),
        ("total_cost", "إجمالي تكلفة الرحلة (محسوب)"),
        ("net_profit", "صافي ربح الرحلة (محسوب)"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
        ("updated_at", "تاريخ آخر تعديل"),
    ],
    "Maintenance": [
        ("id", "المعرف"),
        ("vehicle_id", "معرف العربية"),
        ("date", "التاريخ"),
        ("expense_type", "نوع المصروف"),
        ("description", "الوصف"),
        ("amount", "المبلغ"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Customer_Payments": [
        ("id", "المعرف"),
        ("customer_id", "معرف العميل"),
        ("date", "التاريخ"),
        ("amount", "المبلغ"),
        ("method", "طريقة الدفع"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Owner_Payments": [
        ("id", "المعرف"),
        ("owner_id", "معرف صاحب العربية"),
        ("date", "التاريخ"),
        ("amount", "المبلغ"),
        ("method", "طريقة الدفع"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Driver_Payments": [
        ("id", "المعرف"),
        ("driver_id", "معرف السواق"),
        ("date", "التاريخ"),
        ("amount", "المبلغ"),
        ("method", "طريقة الدفع"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Expenses": [
        ("id", "المعرف"),
        ("date", "التاريخ"),
        ("expense_type", "نوع المصروف"),
        ("description", "الوصف"),
        ("amount", "المبلغ"),
        ("notes", "ملاحظات"),
        ("is_deleted", "محذوف"),
        ("created_at", "تاريخ الإضافة"),
    ],
    "Monthly_Summary": [
        ("month", "الشهر"),
        ("total_revenue", "إجمالي الإيرادات"),
        ("total_expenses", "إجمالي المصروفات"),
        ("net_profit", "صافي الربح"),
        ("collected", "المحصل من العملاء"),
        ("customer_dues", "مستحق على العملاء"),
        ("owner_dues", "مستحق لأصحاب العربيات"),
        ("driver_dues", "مستحق للسواقين"),
        ("trip_count", "عدد الرحلات"),
        ("generated_at", "وقت آخر تحديث"),
    ],
}

HEADER_FILL = PatternFill(start_color="17233D", end_color="17233D", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)


# ---------------------------------------------------------------------------
# إنشاء وفتح وحفظ الملف
# ---------------------------------------------------------------------------
def create_new_workbook():
    """ينشئ Workbook جديد فاضي بكل الشيتات والعناوين المطلوبة فقط (بدون بيانات)."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, columns in SHEETS.items():
        ws = wb.create_sheet(sheet_name)
        headers = [c[1] for c in columns]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"
        # عرض أعمدة معقول بشكل تقريبي
        for col_idx, header in enumerate(headers, start=1):
            width = max(14, min(28, len(header) + 4))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
    return wb


def init_db(force=False):
    """ينشئ ملف الإكسل لو مش موجود. لو force=True يعمل نسخة احتياطية وينشئ ملف جديد فاضي."""
    if DB_PATH.exists() and not force:
        return False
    if DB_PATH.exists() and force:
        create_backup(label="قبل_إعادة_التهيئة")
    wb = create_new_workbook()
    wb.save(DB_PATH)
    return True


@contextmanager
def transaction(timeout=15):
    """
    Context manager لأي عملية تعديل (إضافة/تحديث/حذف).
    يمسك Lock طول عملية القراءة + التعديل + الحفظ عشان يمنع تلف الملف
    أو ضياع تعديلات لو حصل أكتر من طلب في نفس اللحظة.
    الحفظ يتم بطريقة Atomic: بيكتب في ملف مؤقت ثم يستبدل الملف الأصلي دفعة واحدة.
    """
    LOCK_PATH.parent.mkdir(exist_ok=True, parents=True)
    lock = FileLock(str(LOCK_PATH), timeout=timeout)
    with lock:
        if not DB_PATH.exists():
            wb = create_new_workbook()
        else:
            wb = load_workbook(DB_PATH)
        yield wb
        tmp_path = str(DB_PATH) + ".tmp"
        wb.save(tmp_path)
        os.replace(tmp_path, DB_PATH)  # عملية استبدال ذرية (atomic)


@contextmanager
def read_only():
    """فتح الملف للقراءة فقط (أسرع، بدون قفل، آمن لأن الكتابة دايمًا Atomic)."""
    if not DB_PATH.exists():
        wb = create_new_workbook()
    else:
        wb = load_workbook(DB_PATH, data_only=True)
    yield wb


# ---------------------------------------------------------------------------
# عمليات عامة (Generic CRUD) تعمل على أي شيت بالاعتماد على SHEETS
# ---------------------------------------------------------------------------
def _keys(sheet_name):
    return [c[0] for c in SHEETS[sheet_name]]


def _is_deleted_value(v):
    return v in (True, "True", "true", 1, "1")


def get_all(wb, sheet_name, include_deleted=False):
    """يرجع كل الصفوف كـ list of dict (بالمفاتيح الإنجليزية)."""
    ws = wb[sheet_name]
    keys = _keys(sheet_name)
    has_deleted_col = "is_deleted" in keys
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[0] == "":
            continue
        d = dict(zip(keys, row))
        if has_deleted_col and not include_deleted and _is_deleted_value(d.get("is_deleted")):
            continue
        out.append(d)
    return out


def get_by_id(wb, sheet_name, row_id, include_deleted=True):
    for d in get_all(wb, sheet_name, include_deleted=include_deleted):
        if str(d.get("id")) == str(row_id):
            return d
    return None


def get_next_id(wb, sheet_name, start=1):
    ws = wb[sheet_name]
    max_id = start - 1
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[0] == "":
            continue
        try:
            rid = int(row[0])
            max_id = max(max_id, rid)
        except (ValueError, TypeError):
            continue
    return max_id + 1


def append_row(wb, sheet_name, data):
    ws = wb[sheet_name]
    keys = _keys(sheet_name)
    row = [data.get(k, "") for k in keys]
    ws.append(row)


def _find_row_idx(ws, row_id):
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is not None and str(row[0]) == str(row_id):
            return idx
    return None


def update_row(wb, sheet_name, row_id, updates):
    ws = wb[sheet_name]
    idx = _find_row_idx(ws, row_id)
    if idx is None:
        return False
    keys = _keys(sheet_name)
    for key, value in updates.items():
        if key in keys:
            col = keys.index(key) + 1
            ws.cell(row=idx, column=col, value=value)
    return True


def soft_delete(wb, sheet_name, row_id):
    if "is_deleted" not in _keys(sheet_name):
        return False
    return update_row(wb, sheet_name, row_id, {"is_deleted": True})


def clear_sheet_data(wb, sheet_name):
    """يمسح كل الصفوف ما عدا صف العناوين (يُستخدم لتحديث Monthly_Summary وإعادة التهيئة)."""
    ws = wb[sheet_name]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)


# ---------------------------------------------------------------------------
# النسخ الاحتياطي
# ---------------------------------------------------------------------------
def create_backup(label=None):
    if not DB_PATH.exists():
        return None
    ts = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    suffix = f"_{label}" if label else ""
    name = f"backup_{ts}{suffix}.xlsx"
    dest = BACKUP_DIR / name
    shutil.copy2(DB_PATH, dest)
    return dest


def list_backups():
    if not BACKUP_DIR.exists():
        return []
    files = sorted(BACKUP_DIR.glob("backup_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    return files
